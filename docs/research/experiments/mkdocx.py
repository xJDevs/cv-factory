import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_TAB_ALIGNMENT
d = Document()
d.core_properties.title = "Jane Doe - Resume"; d.core_properties.author = "Jane Doe"
d.sections[0].header.paragraphs[0].text = "HEADER: jane.doe@example.com | +1 555 010 2030"
d.add_paragraph("Jane Doe", style="Title")
d.add_heading("Experience", level=1)
p = d.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
p.add_run("Senior Software Engineer, Acme Financial Inc.").bold = True
p.add_run("\tJan 2020 – Present")
d.add_paragraph("Efficient office workflow; financial fluency.", style="List Bullet")
d.add_heading("Skills", level=1)
t = d.add_table(rows=2, cols=2)
t.cell(0,0).text = "LANGUAGES"; t.cell(0,1).text = "PROJECTS"
t.cell(1,0).text = "Python, SQL"; t.cell(1,1).text = "Streaming pipeline processing two billion events per day."
d.save(sys.argv[1])
