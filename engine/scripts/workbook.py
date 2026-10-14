# Copyright 2026 Johel Gomez Rojas. Part of CV Factory. Licensed under the Apache License, Version 2.0 (see LICENSE and NOTICE at the project root). Do not remove this notice.
"""Spreadsheet source: the person's data in one Excel workbook instead of Notion.

Usage:
  python3 workbook.py create <spec.json> <out.xlsx>              empty, ready-to-fill workbook
  python3 workbook.py read   <spec.json> <file.xlsx> <read-folder>
  python3 workbook.py upsert <spec.json> <file.xlsx> <rows.json>    add or update rows
  python3 workbook.py schema <spec.json> <out-schema.json>

"create" writes a first sheet with the spec's instructions and one sheet per table: the ID
column first, a bold, frozen header row with filters, drop-down lists for select and checkbox
columns, and a note on each column title.

"read" writes experience.json, achievements.json and skills.json into <read-folder>, in the
format build_snapshot.py expects, with an explicit "id" per row (the ID cell):
{"database", "rows_in_source", "rows": [{"id", "properties", "content"}]}
Every column except ID goes into "properties" as a string (checkboxes as "__YES__"/"__NO__",
dates as YYYY-MM-DD); "content" is the table's "content_column" (or ""). A row with data and
no ID, a duplicate ID, an ID without the table's prefix, a missing column or a checkbox that is
neither of the two labels stops the script. A row that only has an ID counts as empty.

"upsert" adds or updates rows without opening Excel. rows.json is
{"<table key>": [{"ID": "A7", "<column name>": value, ...}, ...]}  (the ID column's name in the spec)
A row whose ID is already in the sheet gets only the given columns updated; any other row is
written in the first empty row. Checkboxes take true/false, "yes"/"no" or the spec's labels and
are written as the labels; dates take "YYYY-MM-DD" or "YYYY-MM" (the 1st of the month) and are
written as real dates; numbers as numbers; everything else as text; null or "" empties a cell.
Everything is checked first (table, columns, ID prefix, select options, duplicate IDs): if
anything is wrong nothing is written. Drop-down lists, notes and formatting are kept.

"schema" writes the schema.json that build_snapshot.py takes (text fields and verified field
of the three inventory tables), so the next step is:
  python3 build_snapshot.py <read-folder> <out/snapshot.json> --schema <out-schema.json> --hub-read
(--hub-read only once profile.json was read too.)

The spec is passed by path (for example setup/workspace/workbook-spec.json in cv-factory, or
configuracion/espacio/workbook-spec.json in buscador-de-trabajo) and is:
{"language": "en", "checkbox_labels": ["Yes", "No"],
 "readme": {"sheet": "README", "title": "...", "lines": ["...", "## A heading", ...]},
 "messages": {"id", "select", "date", "number"},   (optional: drop-down error texts)
 "tables": [{"key": "experience" | "achievements" | "skills" | <other>, "sheet": "...",
             "id_prefix": "E", "content_column": "Summary" | null,
             "text_fields": [...], "verified_field": "..." | null,   (inventory tables only)
             "columns": [{"name": "ID", "type": "id"},
                         {"name": "...", "type": "text" | "long_text" | "select" | "checkbox"
                                          | "date" | "number", "options": [...], "note": "...",
                          "width": N}]}]}
"""
import argparse
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import platform_tools

INVENTORY = ("experience", "achievements", "skills")
TYPES = ("id", "text", "long_text", "select", "checkbox", "date", "number")
YES, NO = "__YES__", "__NO__"
LAST_ROW = 1000          # drop-downs and formats cover rows 2..LAST_ROW
HEADER_FILL = "DCE6F1"
MESSAGES = {"id": "The ID must start with {prefix} and be unique in this sheet ({prefix}1, {prefix}2, ...).",
            "select": "Pick one of the options in the list.",
            "date": "Type a date, for example 2023-02-01.",
            "number": "Type a number."}
SHEET_NAME_BAD = re.compile(r"[\[\]:*?/\\]")


def fail(message):
    sys.exit(f"ERROR: {message}")


# ---------- spec ----------

def load_spec(path):
    try:
        spec = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as e:
        fail(f"the spec {path} could not be read ({e})")
    problems = []
    labels = spec.get("checkbox_labels")
    if not (isinstance(labels, list) and len(labels) == 2 and all(isinstance(x, str) and x for x in labels)):
        problems.append('"checkbox_labels" must be two labels, like ["Yes", "No"]')
    readme = spec.get("readme", {})
    if not readme.get("sheet") or not isinstance(readme.get("lines"), list):
        problems.append('"readme" needs "sheet" and "lines"')
    keys, sheets, prefixes = set(), {readme.get("sheet")}, set()
    for t in spec.get("tables", []):
        where = f'table "{t.get("sheet") or t.get("key")}"'
        for field in ("key", "sheet", "id_prefix", "columns"):
            if not t.get(field):
                problems.append(f'{where}: missing "{field}"')
        if not t.get("sheet") or not t.get("columns"):
            continue
        if len(t["sheet"]) > 31 or SHEET_NAME_BAD.search(t["sheet"]):
            problems.append(f'{where}: a sheet name has at most 31 characters and none of [ ] : * ? / \\')
        for value, seen, what in ((t["key"], keys, "key"), (t["sheet"], sheets, "sheet"),
                                  (t["id_prefix"], prefixes, "id_prefix")):
            if value in seen:
                problems.append(f'{where}: duplicate {what} "{value}"')
            seen.add(value)
        names = [c.get("name") for c in t["columns"]]
        if t["columns"][0].get("type") != "id":
            problems.append(f'{where}: the first column must be the ID column (type "id")')
        if len(set(names)) != len(names):
            problems.append(f"{where}: duplicate column names")
        for c in t["columns"]:
            if c.get("type") not in TYPES:
                problems.append(f'{where}, column "{c.get("name")}": type must be one of {", ".join(TYPES)}')
            if c.get("type") == "select":
                options = c.get("options") or []
                if not options:
                    problems.append(f'{where}, column "{c["name"]}": a select column needs "options"')
                elif any("," in o or '"' in o for o in options) or len(",".join(options)) > 250:
                    problems.append(f'{where}, column "{c["name"]}": options cannot have commas or quotes, '
                                    'and all of them together must fit in 250 characters')
        by_name = {c.get("name"): c for c in t["columns"]}
        if t.get("content_column") and t["content_column"] not in by_name:
            problems.append(f'{where}: content_column "{t["content_column"]}" is not a column')
        if t["key"] in INVENTORY:
            for f in t.get("text_fields") or []:
                if f not in by_name:
                    problems.append(f'{where}: text field "{f}" is not a column')
            if not t.get("text_fields"):
                problems.append(f'{where}: an inventory table needs "text_fields"')
            vf = t.get("verified_field")
            if vf and by_name.get(vf, {}).get("type") != "checkbox":
                problems.append(f'{where}: verified_field "{vf}" must be a checkbox column')
    missing = [k for k in INVENTORY if k not in keys]
    if missing:
        problems.append("the spec has no table for: " + ", ".join(missing))
    if problems:
        fail("the spec is not valid:\n  " + "\n  ".join(problems))
    return spec


def inventory_tables(spec):
    return {t["key"]: t for t in spec["tables"] if t["key"] in INVENTORY}


# ---------- create ----------

def column_width(col, labels):
    if col.get("width"):
        return col["width"]
    kind = col["type"]
    base = {"id": 8, "text": 28, "long_text": 50, "date": 13, "number": 11,
            "checkbox": max(len(x) for x in labels) + 6}.get(kind)
    if kind == "select":
        base = min(max(len(o) for o in col["options"]) + 5, 40)
    return max(base, min(len(col["name"]) + 4, 40))


def create(spec_path, out):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    spec = load_spec(spec_path)
    yes, no = spec["checkbox_labels"]
    messages = {**MESSAGES, **spec.get("messages", {})}
    wb = Workbook()
    readme = wb.active
    readme.title = spec["readme"]["sheet"]
    readme.column_dimensions["A"].width = 110
    row = 1
    if spec["readme"].get("title"):
        readme.cell(row=1, column=1, value=spec["readme"]["title"]).font = Font(bold=True, size=16)
        row = 3
    for line in spec["readme"]["lines"]:
        cell = readme.cell(row=row, column=1, value=line or None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("## "):
            cell.value = line[3:]
            cell.font = Font(bold=True, size=12)
        row += 1
    readme.sheet_view.showGridLines = False

    for t in spec["tables"]:
        ws = wb.create_sheet(t["sheet"])
        cols = t["columns"]
        for i, col in enumerate(cols, start=1):
            letter = get_column_letter(i)
            cell = ws.cell(row=1, column=i, value=col["name"])
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", start_color=HEADER_FILL)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            note = col.get("note")
            if col["type"] == "id":
                note = note or f'{t["id_prefix"]}1, {t["id_prefix"]}2, ...'
            if note:
                comment = Comment(note, "CV Factory")
                comment.width, comment.height = 300, 110
                cell.comment = comment
            ws.column_dimensions[letter].width = column_width(col, (yes, no))
            cells = f"{letter}2:{letter}{LAST_ROW}"
            dv = None
            if col["type"] == "id":
                p = t["id_prefix"]
                dv = DataValidation(type="custom", allow_blank=True,
                                    formula1=f'AND(LEFT({letter}2,{len(p)})="{p}",'
                                             f'COUNTIF(${letter}$2:${letter}${LAST_ROW},{letter}2)=1)')
                dv.error = messages["id"].format(prefix=p)
            elif col["type"] in ("select", "checkbox"):
                options = col["options"] if col["type"] == "select" else [yes, no]
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                dv.error = messages["select"]
                if col["type"] == "select":
                    dv.errorStyle = "warning"
            elif col["type"] == "date":
                dv = DataValidation(type="date", operator="between", allow_blank=True,
                                    formula1="DATE(1950,1,1)", formula2="DATE(2100,12,31)")
                dv.errorStyle = "warning"
                dv.error = messages["date"]
                for r in range(2, LAST_ROW + 1):
                    ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
            elif col["type"] == "number":
                dv = DataValidation(type="decimal", operator="greaterThanOrEqual", allow_blank=True,
                                    formula1="0")
                dv.errorStyle = "warning"
                dv.error = messages["number"]
            elif col["type"] == "long_text":
                for r in range(2, LAST_ROW + 1):
                    ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
            if dv is not None:
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(cells)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Workbook: wrote {out} ({len(spec['tables'])} tables + {spec['readme']['sheet']})")


# ---------- read ----------

def fold(text):
    return "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)).casefold()


def as_string(value, kind, labels, where):
    from openpyxl.utils.datetime import from_excel

    if kind == "checkbox":
        if value is None or (isinstance(value, str) and not value.strip()):
            return NO
        if isinstance(value, bool):
            return YES if value else NO
        text = fold(str(value).strip())
        if text in (fold(labels[0]), "true"):
            return YES
        if text in (fold(labels[1]), "false"):
            return NO
        fail(f'{where}: "{value}" is not "{labels[0]}" or "{labels[1]}"')
    if value is None:
        return ""
    if kind == "date" and isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value)
    if isinstance(value, datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(wb, table, labels, file_name):
    sheet = table["sheet"]
    if sheet not in wb.sheetnames:
        fail(f'{file_name}: the sheet "{sheet}" is missing (sheets must keep their names)')
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else [])]
    position = {}
    for i, name in enumerate(header):
        if name:
            if name in position:
                fail(f'"{sheet}": the column "{name}" appears twice')
            position[name] = i
    spec_cols = table["columns"]
    missing = [c["name"] for c in spec_cols if c["name"] not in position]
    if missing:
        fail(f'"{sheet}": missing column(s) {", ".join(missing)} (column titles must keep their names)')
    kinds = {c["name"]: c["type"] for c in spec_cols}
    id_name = spec_cols[0]["name"]
    extra = [n for n in position if n not in kinds]
    out, no_id, seen, bad_prefix = [], [], {}, []
    for number, values in enumerate(rows[1:], start=2):
        cells = {name: (values[i] if i < len(values) else None) for name, i in position.items()}
        filled = {n: v for n, v in cells.items() if v is not None and str(v).strip() != ""}
        rid = as_string(cells[id_name], "id", labels, "")
        others = [n for n in filled if n != id_name]
        if not others:
            continue                    # empty row, or only an ID typed in advance
        if not rid:
            no_id.append(number)
            continue
        if rid in seen:
            fail(f'"{sheet}": the ID {rid} is used twice (rows {seen[rid]} and {number}); every ID must be unique')
        seen[rid] = number
        if not rid.startswith(table["id_prefix"]):
            bad_prefix.append(f"{rid} (row {number})")
        props = {}
        for name in [c["name"] for c in spec_cols[1:]] + extra:
            props[name] = as_string(cells[name], kinds.get(name, "text"), labels,
                                    f'"{sheet}", row {number}, column "{name}"')
        content = props.get(table.get("content_column") or "", "")
        out.append({"id": rid, "properties": props, "content": content})
    if no_id:
        fail(f'"{sheet}": row(s) {", ".join(map(str, no_id))} have data but no ID. '
             f'Give each one a unique ID ({table["id_prefix"]}1, {table["id_prefix"]}2, ...): the CV cites rows by ID')
    if bad_prefix:
        fail(f'"{sheet}": IDs must start with {table["id_prefix"]}: ' + ", ".join(bad_prefix))
    return {"database": sheet, "rows_in_source": len(out), "rows": out}


def read(spec_path, xlsx, folder):
    from openpyxl import load_workbook

    spec = load_spec(spec_path)
    try:
        wb = load_workbook(xlsx, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 (any reason the file cannot be opened)
        fail(f"{xlsx} could not be opened as an Excel workbook ({e})")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    counts, all_ids = {}, {}
    results = {}
    for key, table in inventory_tables(spec).items():
        d = read_table(wb, table, spec["checkbox_labels"], Path(xlsx).name)
        for row in d["rows"]:
            if row["id"] in all_ids:
                fail(f'the ID {row["id"]} is in "{all_ids[row["id"]]}" and in "{table["sheet"]}"')
            all_ids[row["id"]] = table["sheet"]
        results[key] = d
        counts[key] = d["rows_in_source"]
    wb.close()
    for key, d in results.items():
        (folder / f"{key}.json").write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8", newline="\n")
    print(f"Workbook: read {counts} → {folder}")


# ---------- upsert ----------

def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def cell_value(value, col, labels, where, problems):
    """The value to write for one cell, or problems.append(...) and None."""
    kind = col["type"]
    if blank(value):
        return None
    if isinstance(value, (dict, list)):
        problems.append(f"{where}: a value must be text, a number, true/false or null")
        return None
    if kind == "checkbox":
        if isinstance(value, bool):
            return labels[0] if value else labels[1]
        text = fold(str(value).strip())
        if text in (fold(labels[0]), "yes", "true"):
            return labels[0]
        if text in (fold(labels[1]), "no", "false"):
            return labels[1]
        problems.append(f'{where}: "{value}" is not true/false, yes/no, "{labels[0]}" or "{labels[1]}"')
        return None
    if kind == "select":
        match = [o for o in col["options"] if fold(o) == fold(str(value).strip())]
        if not match:
            problems.append(f'{where}: "{value}" is not one of the options: {", ".join(col["options"])}')
            return None
        return match[0]
    if kind == "date":
        m = re.fullmatch(r"(\d{4})-(\d{2})(?:-(\d{2}))?", str(value).strip())
        try:
            if not m or isinstance(value, bool):
                raise ValueError
            year, month = int(m.group(1)), int(m.group(2))
            return date(year, month, int(m.group(3) or 1))
        except ValueError:
            problems.append(f'{where}: "{value}" is not a date (use YYYY-MM-DD or YYYY-MM)')
            return None
    if kind == "number":
        if isinstance(value, bool):
            problems.append(f'{where}: {value} is not a number')
            return None
        try:
            number = float(str(value).strip().replace(",", "")) if isinstance(value, str) else float(value)
        except ValueError:
            problems.append(f'{where}: "{value}" is not a number')
            return None
        return int(number) if number.is_integer() else number
    if isinstance(value, bool):
        problems.append(f"{where}: true/false only fits a checkbox column")
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def sheet_rows(ws, table, problems):
    """(column number per spec column, {ID: row number}, used row numbers) of one sheet."""
    sheet = table["sheet"]
    header = {}
    for i, cell in enumerate(ws[1], start=1):
        name = str(cell.value).strip() if cell.value is not None else ""
        if name:
            if name in header:
                problems.append(f'"{sheet}": the column "{name}" appears twice')
            header[name] = i
    missing = [c["name"] for c in table["columns"] if c["name"] not in header]
    if missing:
        problems.append(f'"{sheet}": missing column(s) {", ".join(missing)} (column titles must keep their names)')
        return None, {}, set()
    id_col = header[table["columns"][0]["name"]]
    ids, used = {}, set()
    for number in range(2, ws.max_row + 1):
        values = [ws.cell(row=number, column=i).value for i in header.values()]
        if all(blank(v) for v in values):
            continue
        used.add(number)
        rid = ws.cell(row=number, column=id_col).value
        if blank(rid):
            continue
        rid = as_string(rid, "id", ("", ""), "")
        if rid in ids:
            problems.append(f'"{sheet}": the ID {rid} is already used twice (rows {ids[rid]} and {number}); fix the sheet first')
        ids[rid] = number
    return header, ids, used


def upsert(spec_path, xlsx, rows_path):
    import os
    import tempfile
    from openpyxl import load_workbook

    spec = load_spec(spec_path)
    labels = spec["checkbox_labels"]
    tables = {t["key"]: t for t in spec["tables"]}
    try:
        wanted = json.loads(Path(rows_path).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as e:
        fail(f"{rows_path} could not be read ({e})")
    if not isinstance(wanted, dict) or not wanted:
        fail(f'{rows_path} must be an object like {{"achievements": [{{"ID": "A7", ...}}]}}')
    xlsx = Path(xlsx)
    for lock in (xlsx.with_name("~$" + xlsx.name), xlsx.with_name(f".~lock.{xlsx.name}#")):
        if lock.exists():
            fail(f"{xlsx.name} seems to be open in Excel or LibreOffice ({lock.name} exists): "
                 "close it and try again, or the changes would be lost")
    try:
        wb = load_workbook(xlsx)
    except Exception as e:  # noqa: BLE001 (any reason the file cannot be opened)
        fail(f"{xlsx} could not be opened as an Excel workbook ({e})")

    problems, plan, seen_ids = [], [], {}
    for key, rows in wanted.items():
        table = tables.get(key)
        if table is None:
            problems.append(f'unknown table "{key}" (the spec has: {", ".join(tables)})')
            continue
        sheet = table["sheet"]
        if not isinstance(rows, list):
            problems.append(f'"{key}": must be a list of rows')
            continue
        if sheet not in wb.sheetnames:
            problems.append(f'{xlsx.name}: the sheet "{sheet}" is missing (sheets must keep their names)')
            continue
        ws = wb[sheet]
        header, existing, used = sheet_rows(ws, table, problems)
        if header is None:
            continue
        columns = {c["name"]: c for c in table["columns"]}
        id_name = table["columns"][0]["name"]
        updated = added = 0
        writes = []
        for n, row in enumerate(rows, start=1):
            where = f'"{key}" row {n}'
            if not isinstance(row, dict):
                problems.append(f"{where}: must be an object with an {id_name} and columns")
                continue
            rid = row.get(id_name)
            rid = "" if blank(rid) else str(rid).strip()
            if not rid:
                problems.append(f'{where}: missing "{id_name}"')
                continue
            where = f'"{key}" {rid}'
            if not rid.startswith(table["id_prefix"]):
                problems.append(f'{where}: the ID must start with {table["id_prefix"]}')
            if rid in seen_ids:
                problems.append(f"{where}: the ID appears twice in {Path(rows_path).name} ({seen_ids[rid]})")
            seen_ids[rid] = f'"{key}"'
            unknown = [c for c in row if c not in columns]
            if unknown:
                problems.append(f'{where}: unknown column(s) {", ".join(unknown)} '
                                f'(the sheet "{sheet}" has: {", ".join(columns)})')
            values = {}
            for name, value in row.items():
                if name in columns and name != id_name:
                    values[name] = cell_value(value, columns[name], labels, f'{where}, column "{name}"', problems)
            if rid in existing:
                updated += 1
                writes.append((existing[rid], rid, values, False))
            else:
                if all(v is None for v in values.values()):
                    problems.append(f"{where}: a new row needs at least one column with data besides {id_name}")
                added += 1
                writes.append((None, rid, values, True))
        plan.append((table, ws, header, used, writes, updated, added))
    if problems:
        fail(f"nothing was written to {xlsx.name}:\n  " + "\n  ".join(problems))

    for table, ws, header, used, writes, _, _ in plan:
        id_name = table["columns"][0]["name"]
        free = 2
        for number, rid, values, new in writes:
            if new:
                while free in used:
                    free += 1
                number = free
                used.add(number)
                ws.cell(row=number, column=header[id_name], value=rid)
            for name, value in values.items():
                cell = ws.cell(row=number, column=header[name], value=value)
                if isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
    for t in spec["tables"]:            # openpyxl forgets the size of the notes when it loads them
        if t["sheet"] in wb.sheetnames:
            for cell in wb[t["sheet"]][1]:
                if cell.comment is not None:
                    cell.comment.width, cell.comment.height = 300, 110
    handle, temp = tempfile.mkstemp(suffix=".xlsx", dir=xlsx.parent)
    os.close(handle)
    try:
        wb.save(temp)
        os.replace(temp, xlsx)
    except PermissionError:   # Windows: Excel keeps the file locked while it is open
        fail(f"{xlsx.name} is open in Excel or LibreOffice (it could not be replaced): "
             "close it and try again; nothing was written")
    finally:
        Path(temp).unlink(missing_ok=True)
    for table, _, _, _, _, updated, added in plan:
        print(f'{table["sheet"]}: {updated} updated, {added} added')


# ---------- schema ----------

def schema(spec_path, out):
    spec = load_spec(spec_path)
    result = {key: {"text_fields": t["text_fields"], "verified_field": t.get("verified_field")}
              for key, t in inventory_tables(spec).items()}
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    Path(out).write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8", newline="\n")
    print(f"Workbook: wrote the snapshot schema to {out}")


def main():
    ap = argparse.ArgumentParser(description="Create, read or update the person's data workbook (spreadsheet source).")
    sub = ap.add_subparsers(dest="command", required=True)
    c = sub.add_parser("create", help="create an empty workbook from a spec")
    c.add_argument("spec")
    c.add_argument("out", help="path of the .xlsx to write")
    r = sub.add_parser("read", help="read a filled workbook into a read folder for build_snapshot.py")
    r.add_argument("spec")
    r.add_argument("xlsx")
    r.add_argument("folder", help="read folder (experience.json, achievements.json, skills.json)")
    u = sub.add_parser("upsert", help="add or update rows of a workbook from a rows.json")
    u.add_argument("spec")
    u.add_argument("xlsx")
    u.add_argument("rows", help='rows.json: {"<table key>": [{"ID": "A7", "<column>": value, ...}]}')
    s = sub.add_parser("schema", help="write the schema.json for build_snapshot.py")
    s.add_argument("spec")
    s.add_argument("out", help="path of the schema.json to write")
    a = ap.parse_args()
    if a.command == "create":
        create(a.spec, a.out)
    elif a.command == "read":
        read(a.spec, a.xlsx, a.folder)
    elif a.command == "upsert":
        upsert(a.spec, a.xlsx, a.rows)
    else:
        schema(a.spec, a.out)
    return 0


if __name__ == "__main__":
    platform_tools.utf8_console()
    sys.exit(main())
