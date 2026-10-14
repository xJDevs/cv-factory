# Copyright 2026 Johel Gomez Rojas. Part of CV Factory. Licensed under the Apache License, Version 2.0 (see LICENSE and NOTICE at the project root). Do not remove this notice.
"""View → DOCX with python-docx: paragraphs only, no tables, text boxes or headers.

The bullet is defined by hand as '•' in Calibri: Word's default bullet is U+F0B7
(Symbol) and comes out as garbage when the text is extracted (research doc 05, §5.3).
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Calibri"  # Carlito is its metric equivalent on Linux
SIZES = {"letter": (Inches(8.5), Inches(11)), "a4": (Inches(8.27), Inches(11.69))}
NUM_ID = 9


def set_font(style_or_run, size, bold=False):
    f = style_or_run.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, False
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(attr), None)


def prepare_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.space_before = Pt(0)
    for name, size in (("Title", 16), ("Heading 1", 11.5)):
        st = doc.styles[name]
        set_font(st, size, bold=True)
        ppr = st.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)
        st.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 0)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.keep_with_next = True


def bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def define_bullet(doc):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(NUM_ID))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"),
                       ("w:lvlJc", "left")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), value)
        level.append(e)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "240")
    ind.set(qn("w:hanging"), "180")
    ppr.append(ind)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), FONT)
    rf.set(qn("w:hint"), "default")
    rpr.append(rf)
    level.append(rpr)
    abstract.append(level)
    numbering.insert(0, abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(NUM_ID))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(NUM_ID))
    num.append(ref)
    numbering.append(num)


def bullet(doc, text):
    p = doc.add_paragraph(text)
    numpr = OxmlElement("w:numPr")
    for tag, value in (("w:ilvl", "0"), ("w:numId", str(NUM_ID))):
        e = OxmlElement(tag)
        e.set(qn("w:val"), value)
        numpr.append(e)
    p._p.get_or_add_pPr().insert(0, numpr)


def generate(view, folder, base):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = SIZES.get(view["paper"], SIZES["letter"])
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.7))
    prepare_styles(doc)
    define_bullet(doc)
    doc.add_paragraph(view["name"], style="Title")
    doc.add_paragraph(view["headline"])
    doc.add_paragraph(" | ".join(view["contact"]))
    for s in view["sections"]:
        bottom_border(doc.add_paragraph(s["title"], style="Heading 1"))
        if "paragraph" in s:
            doc.add_paragraph(s["paragraph"])
        for l in s.get("labeled", []):
            p = doc.add_paragraph()
            p.add_run(l["label"]).bold = True
            p.add_run(" " + l["text"])
        for e in s.get("entries", []):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = bool(e["bullets"])
            p.paragraph_format.space_before = Pt(4)
            p.add_run(e["bold"]).bold = True
            p.add_run(e["rest"] + (f" | {e['dates']}" if e["dates"] else ""))
            for b in e["bullets"]:
                bullet(doc, b)
    cp = doc.core_properties
    cp.title, cp.author, cp.subject = view["pdf_title"], view["name"], view["headline"]
    cp.last_modified_by, cp.comments, cp.keywords, cp.category = view["name"], "", "", ""
    path = Path(folder) / f"{base}.docx"
    doc.save(path)
    return path
